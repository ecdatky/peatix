#!/usr/bin/env python3
"""Peatix「RIK SQUARE! 2026」の参加者リストをダウンロードし、確認用シート付きxlsxに整形する。
使い方: python3 fetch_peatix_csv.py <出力先ディレクトリ>
環境変数 PEATIX_USER / PEATIX_PASSWORD が必要。"""
import csv, io, os, re, sys, unicodedata
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.styles.colors import Color
from playwright.sync_api import sync_playwright

EVENT = "5227975"
out_dir = sys.argv[1]
os.makedirs(out_dir, exist_ok=True)
with sync_playwright() as p:
    b = p.chromium.launch(executable_path="/opt/pw-browsers/chromium")
    pg = b.new_page()
    pg.goto("https://peatix.com/signin", wait_until="domcontentloaded")
    pg.wait_for_timeout(3000)
    pg.fill("input[name=username]", os.environ["PEATIX_USER"])
    pg.click('button:has-text("Next")'); pg.wait_for_timeout(3000)
    pg.fill("input[type=password]", os.environ["PEATIX_PASSWORD"])
    pg.keyboard.press("Enter"); pg.wait_for_timeout(6000)
    pg.goto(f"https://peatix.com/event/{EVENT}/list_sales", wait_until="domcontentloaded")
    pg.wait_for_timeout(4000)
    with pg.expect_download(timeout=60000) as d:
        pg.click("text=参加者リスト(CSV)")
    path = os.path.join(out_dir, "participants.csv")
    d.value.save_as(path)

# Peatix標準はUTF-16LE・タブ区切り。
text = open(path, encoding="utf-16").read()
rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text, newline=""), delimiter="\t")]
header, data = rows[0], [r for r in rows[1:] if any(r)]

now = datetime.now(ZoneInfo("Asia/Tokyo"))
wb = Workbook()

# 確認用シート: チケット番号以降の列だけ残し、郵便番号をハイフン入りに、ヘッダーに色を付ける
ws = wb.active
ws.title = "確認用"
start = header.index("チケット番号")
ci_zip = header.index("郵便番号") - start


def fmt_zip(v):
    d = re.sub(r"\D", "", unicodedata.normalize("NFKC", v))
    return f"{d[:3]}-{d[3:]}" if len(d) == 7 else v


ws.append(header[start:])
for r in data:
    row = r[start:]
    row[ci_zip] = fmt_zip(row[ci_zip])
    ws.append(row)
fill = PatternFill("solid", fgColor=Color(theme=3, tint=0.8999908444471572))
for c in ws[1]:
    c.fill = fill
for col in ws.columns:
    w = max(len(str(c.value or "")) * 2 if ord(str(c.value or " ")[0]) > 255 else len(str(c.value or "")) for c in col)
    ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 10), 50)
for row in ws.iter_rows():
    for c in row:
        c.alignment = Alignment(vertical="center")
ws.freeze_panes = "A2"

# 元データ(CSVそのまま)シート
raw = wb.create_sheet(f"participants{now:%Y%m%d}")
int_cols = {header.index(h) for h in ("注文番号", "会員 ID", "枚数")}
raw.append(header)
for r in data:
    raw.append([int(v) if i in int_cols and v.isdigit() else v for i, v in enumerate(r)])

out = os.path.join(out_dir, f"participants_{now:%Y%m%d_%H%M}.xlsx")
wb.save(out)
print(out)
