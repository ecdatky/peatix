#!/usr/bin/env python3
"""【参加者リスト】RIK SQUARE!2026.xlsx の初期版(2025年の中身を消した新しい器)を作る。初回の1回だけ使う。
使い方: python3 build_master_template.py <fetch_peatix_csv.py の出力ディレクトリ> <出力xlsx>
シート「参加申込リスト」に Excel テーブル「参加申込」を置き、現時点の参加者を入れておく。
以降は Power Automate が「チケット番号」をキーに、新しい人は行追加、既存の人は Peatix 由来の列だけ更新する。
「講師」以降の列は手入力用で、自動更新では触らない。"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

AUTO = ["完了時刻", "担当営業所", "会社名", "申込ご担当者", "ヨミ", "部署・役職", "都道府県", "郵便番号", "住所", "電話番号",
        "メールアドレス", "ご参加数", "懇親会", "業種", "業種（その他）", "ステータス", "チケット番号"]
MANUAL = ["講師", "卒業生", "備考欄", "新規", "ユーザー", "取引先種類", "情報・商談メモ ※複数申込の場合は話をした方の名前も記載"]
WIDTH = {"完了時刻": 17, "担当営業所": 11, "会社名": 30, "申込ご担当者": 14, "ヨミ": 16, "部署・役職": 20, "都道府県": 10,
         "郵便番号": 10, "住所": 36, "電話番号": 15, "メールアドレス": 30, "ご参加数": 9, "懇親会": 11, "業種": 22,
         "業種（その他）": 20, "ステータス": 11, "チケット番号": 12, "備考欄": 24, "取引先種類": 12, MANUAL[-1]: 40}

rows = json.load(open(os.path.join(sys.argv[1], "master_rows.json"), encoding="utf-8"))
cols = AUTO + MANUAL
n = max(len(rows), 1)  # テーブルは最低1行必要
last = 2 + n
font = Font(name="游ゴシック", size=11)

wb = Workbook()
ws = wb.active
ws.title = "参加申込リスト"
ws["A1"] = "RIK SQUARE!2026 お申込状況"
ws["A1"].font = Font(name="游ゴシック", size=14, bold=True)
# 集計はテーブル外の固定範囲で数える(行追加でテーブルが伸びても届くよう広めに取る)
st, kon = "P", "M"  # ステータス列, 懇親会列
# 社内(会社名にユニマットリックを含む)の申込は数えない
rik = 'C3:C2000,"<>*ユニマットリック*"'
ws["E1"], ws["F1"] = "申込者数", f'=COUNTIFS({st}3:{st}2000,"<>",{st}3:{st}2000,"<>*キャンセル*",{rik})'
ws["G1"], ws["H1"] = "懇親会", f'=COUNTIFS({kon}3:{kon}2000,"参加する",{st}3:{st}2000,"<>*キャンセル*",{rik})'
ws["J1"] = "キャンセル"
ws["K1"] = f'=COUNTIFS({st}3:{st}2000,"*キャンセル*",{rik})'
for c in ("E1", "F1", "G1", "H1", "J1", "K1"):
    ws[c].font = Font(name="游ゴシック", size=11, bold=True)

for j, h in enumerate(cols, 1):
    ws.cell(2, j, h)
for i, r in enumerate(rows, 3):
    for j, h in enumerate(AUTO, 1):
        ws.cell(i, j, r.get(h, ""))
for row in ws.iter_rows(min_row=2, max_row=last, max_col=len(cols)):
    for c in row:
        c.font = font
        c.alignment = Alignment(vertical="center", wrap_text=c.row == 2)
manual_fill = PatternFill("solid", fgColor="FFF2CC")  # 手入力列の見出しは黄色で区別
for j in range(len(AUTO) + 1, len(cols) + 1):
    ws.cell(2, j).fill = manual_fill
    ws.cell(2, j).font = Font(name="游ゴシック", size=11, bold=True, color="000000")
for j, h in enumerate(cols, 1):
    ws.column_dimensions[ws.cell(2, j).column_letter].width = WIDTH.get(h, 9)
ws.row_dimensions[2].height = 32
ws.freeze_panes = "D3"

# 確認用の色付け(行追加で伸びても効くよう広めに取る)
# 社内(ユニマットリック)の行は全体をグレー(最優先) / 会社名: 2行以上ある会社はオレンジ、「個人」は黄色
# メールアドレス: 2行以上あれば赤 / 懇親会: 「参加する」は緑
ws.conditional_formatting.add(f"A3:{ws.cell(2, len(cols)).column_letter}2000", FormulaRule(
    formula=['ISNUMBER(SEARCH("ユニマットリック",$C3))'], fill=PatternFill("solid", bgColor="D9D9D9"), stopIfTrue=True))
ws.conditional_formatting.add("C3:C2000", FormulaRule(
    formula=['AND($C3<>"",$C3<>"個人",COUNTIF($C$3:$C$2000,$C3)>1)'], fill=PatternFill("solid", bgColor="FCE4D6")))
ws.conditional_formatting.add("C3:C2000", FormulaRule(
    formula=['$C3="個人"'], fill=PatternFill("solid", bgColor="FFF2CC")))
ws.conditional_formatting.add("K3:K2000", FormulaRule(
    formula=['AND($K3<>"",COUNTIF($K$3:$K$2000,$K3)>1)'], fill=PatternFill("solid", bgColor="FFC7CE")))
ws.conditional_formatting.add("M3:M2000", FormulaRule(
    formula=['$M3="参加する"'], fill=PatternFill("solid", bgColor="E2EFDA")))

ref = f"A2:{ws.cell(2, len(cols)).column_letter}{last}"
t = Table(displayName="参加申込", ref=ref)
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(t)
wb.save(sys.argv[2])
print(sys.argv[2], f"{len(rows)} rows")
